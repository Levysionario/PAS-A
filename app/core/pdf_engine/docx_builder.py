import os
from docx import Document
from docx.shared import Pt, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image
from app.schemas.prova import ProvaSchema

try:
    import cv2
    import numpy as np
    cv2_disponivel = True
except ImportError:
    cv2_disponivel = False

class ProvaDocxBuilder:
    MAPA_TEMPLATES_DOCX = {
        "acn": "CABEÇALHO DE ACN.docx",
        "atividade": "CABEÇALHO DE ATIVIDADE.docx",
        "bimestral": "CABEÇALHO DE AVALIAÇÃO BIMESTRAL.docx",
        "diagnostica": "CABEÇALHO DE AVALIAÇÃO DIAGNÓSTICA.docx",
        "extra": "CABEÇALHO DE MATERIAL EXTRA.docx",
        "final": "CABEÇALHO DE RECUPERAÇÃO FINAL.docx",
        "guia": "CABEÇALHO DE GUIA DE ESTUDO.docx",
        "mensal": "CABEÇALHO DE AVALIAÇÃO MENSAL.docx",
        "ta": "CABEÇALHO DE TRABALHO AVALIATIVO.docx",
        "trabalho": "CABEÇALHO DE TRABALHO AVALIATIVO.docx"
    }

    def __init__(self, prova: ProvaSchema, output_docx_path: str, tipo_cabecalho: str = "mensal"):
        self.prova = prova
        self.output_docx_path = output_docx_path
        self.tipo_cabecalho = tipo_cabecalho.lower()
        
        nome_template = self.MAPA_TEMPLATES_DOCX.get(self.tipo_cabecalho, "CABEÇALHO DE AVALIAÇÃO MENSAL.docx")
        template_path = os.path.join("storage", "templates", nome_template)

        if not os.path.exists(template_path):
            template_path = os.path.join("storage", "templates", "CABEÇALHO DE AVALIAÇÃO BIMESTRAL.docx")

        if os.path.exists(template_path):
            self.doc = Document(template_path)
            self._preservar_banner_e_limpar_gabarito_estatico()
        else:
            self.doc = Document()
            self._configurar_margens_a4()

    def _configurar_margens_a4(self):
        section = self.doc.sections[0]
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Mm(10)
        section.bottom_margin = Mm(10)
        section.left_margin = Mm(10)
        section.right_margin = Mm(10)

    def _preservar_banner_e_limpar_gabarito_estatico(self):
        """Preserva a primeira imagem encontrada (o banner institucional do topo) e remove apenas as imagens de gabarito subsequentes"""
        desenhos_encontrados = 0
        para_remover = []
        for p in self.doc.paragraphs:
            if p._element.xpath('.//w:drawing'):
                desenhos_encontrados += 1
                # O primeiro elemento de imagem no topo é o banner institucional. Removemos apenas do segundo em diante (gabarito antigo)
                if desenhos_encontrados > 1:
                    para_remover.append(p)

        for p in para_remover:
            p._element.getparent().remove(p._element)

    def _aplicar_estilo_paragrafo(self, p, tamanho_pt=12, negrito=False, espacamento_depois=0, alinhamento=None):
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(espacamento_depois)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.left_indent = Pt(0)
        if alinhamento is not None:
            p.alignment = alinhamento
        
        for run in p.runs:
            run.font.name = 'Arial Narrow'
            run.font.size = Pt(tamanho_pt)
            run.bold = negrito

    def _substituir_placeholders_cabecalho(self):
        disc_nome = self.prova.disciplina.upper() if self.prova.disciplina and self.prova.disciplina != "NÃO INFORMADA" else "DISCIPLINA"
        prof_nome = self.prova.professor if self.prova.professor and self.prova.professor != "Não informado" else "Nome"

        for table in self.doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if "DISCIPLINA" in cell.text or "{{DISCIPLINA}}" in cell.text:
                        cell.text = disc_nome
                        p = cell.paragraphs[0]
                        self._aplicar_estilo_paragrafo(p, tamanho_pt=14, negrito=True, alinhamento=WD_ALIGN_PARAGRAPH.CENTER)
                    elif "Prof." in cell.text or "{{PROFESSOR}}" in cell.text:
                        texto_atual = cell.text
                        if "Prof. Nome" in texto_atual:
                            novo_texto = texto_atual.replace("Prof. Nome", f"Prof. {prof_nome}")
                        elif "{{PROFESSOR}}" in texto_atual:
                            novo_texto = texto_atual.replace("{{PROFESSOR}}", prof_nome)
                        else:
                            novo_texto = texto_atual
                            
                        cell.text = novo_texto
                        p = cell.paragraphs[0]
                        self._aplicar_estilo_paragrafo(p, tamanho_pt=9.5, negrito=True)

    def _recortar_gabarito_opencv(self, caminho_gab: str, num_questoes: int) -> str:
        img_mat = cv2.imread(caminho_gab)
        altura, largura, _ = img_mat.shape

        colunas_necessarias = (num_questoes + 9) // 10
        if colunas_necessarias > 3:
            colunas_necessarias = 3

        gray = cv2.cvtColor(img_mat, cv2.COLOR_BGR2GRAY)
        _, thresh = cv2.threshold(gray, 200, 255, cv2.THRESH_BINARY_INV)

        projecao_vertical = np.sum(thresh, axis=0)
        x_corte_base = int((largura / 3.0) * colunas_necessarias)
        
        janela_busca = projecao_vertical[max(0, x_corte_base - 30):min(largura, x_corte_base + 30)]
        if len(janela_busca) > 0:
            offset_ideal = np.argmin(janela_busca)
            x_corte_exato = (x_corte_base - 30) + offset_ideal
        else:
            x_corte_exato = x_corte_base

        img_recortada = img_mat[:, 0:int(x_corte_exato)]
        
        caminho_temp = os.path.join("storage", "gabarito", "gabarito_cv.png")
        cv2.imwrite(caminho_temp, img_recortada)
        return caminho_temp

    def _recortar_e_inserir_gabarito(self):
        num_questoes = len(self.prova.questoes)
        if num_questoes == 0:
            return

        caminho_gab_png = os.path.join("storage", "gabarito", "gabarito.png")
        caminho_gab_jpg = os.path.join("storage", "gabarito", "gabarito.jpg")
        caminho_gab = caminho_gab_png if os.path.exists(caminho_gab_png) else caminho_gab_jpg

        if not os.path.exists(caminho_gab):
            return

        try:
            colunas_necessarias = (num_questoes + 9) // 10
            if colunas_necessarias > 3:
                colunas_necessarias = 3

            if cv2_disponivel:
                caminho_temp = self._recortar_gabarito_opencv(caminho_gab, num_questoes)
            else:
                img = Image.open(caminho_gab)
                largura_total, altura_total = img.size
                largura_por_bloco = largura_total / 3.0
                largura_corte = int(largura_por_bloco * colunas_necessarias)
                img_recortada = img.crop((0, 0, largura_corte, altura_total))
                caminho_temp = os.path.join("storage", "gabarito", "gabarito_temp.png")
                img_recortada.save(caminho_temp)

            p_gab = self.doc.add_paragraph()
            run_gab = p_gab.add_run()
            
            largura_mm = min(180, colunas_necessarias * 58)
            run_gab.add_picture(caminho_temp, width=Mm(largura_mm))
            self._aplicar_estilo_paragrafo(p_gab, espacamento_depois=8, alinhamento=WD_ALIGN_PARAGRAPH.CENTER)

        except Exception as e:
            print(f"[Aviso] Erro no recorte do gabarito: {e}")

    def build(self) -> str:
        self._substituir_placeholders_cabecalho()
        self._recortar_e_inserir_gabarito()

        for q in self.prova.questoes:
            p_titulo = self.doc.add_paragraph()
            p_titulo.add_run(f"Questão {q.numero}")
            self._aplicar_estilo_paragrafo(p_titulo, tamanho_pt=12, negrito=True, espacamento_depois=0)

            if q.enunciado and q.enunciado.strip():
                p_enunciado = self.doc.add_paragraph()
                p_enunciado.add_run(q.enunciado.strip())
                self._aplicar_estilo_paragrafo(p_enunciado, tamanho_pt=12, negrito=False, espacamento_depois=4)

            if q.imagens:
                for img_path in q.imagens:
                    if os.path.exists(img_path):
                        p_img = self.doc.add_paragraph()
                        run_img = p_img.add_run()
                        try:
                            im = Image.open(img_path)
                            w_px, h_px = im.size
                            largura_desejada = min(140, w_px * 0.264)
                            run_img.add_picture(img_path, width=Mm(largura_desejada))
                        except Exception:
                            run_img.add_picture(img_path, width=Mm(120))
                        self._aplicar_estilo_paragrafo(p_img, espacamento_depois=4, alinhamento=WD_ALIGN_PARAGRAPH.CENTER)

            for alt in q.alternativas:
                p_alt = self.doc.add_paragraph()
                p_alt.add_run(f"{alt.letra.lower()}) {alt.texto.strip()}")
                self._aplicar_estilo_paragrafo(p_alt, tamanho_pt=12, negrito=False, espacamento_depois=2)

            p_qspace = self.doc.add_paragraph()
            self._aplicar_estilo_paragrafo(p_qspace, espacamento_depois=6)

        self.doc.save(self.output_docx_path)
        return self.output_docx_path