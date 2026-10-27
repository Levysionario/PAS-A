import re
import os
import zipfile
from docx import Document
from typing import List, Optional, Dict
from app.schemas.prova import ProvaSchema, QuestaoSchema, AlternativaSchema

try:
    import cv2
    import numpy as np
    cv2_disponivel = True
except ImportError:
    cv2_disponivel = False

try:
    import pytesseract
    from PIL import Image
    pytesseract_disponivel = True
except ImportError:
    pytesseract_disponivel = False

class DocxParser:
    REGEX_NUMERO_QUESTAO = re.compile(
        r'^\s*(?:questã[o|a]|q|item)?\s*[\.\-_\(\)]?\s*(\d+)[\.\-\)\:\s-]\s*(.*)', 
        re.IGNORECASE
    )
    REGEX_ALTERNATIVA = re.compile(
        r'^\s*[\(\[]?\s*([a-eA-E])\s*[\.\)\-\]]\s*(.*)'
    )

    def __init__(self, docx_path: str, media_output_folder: str = "storage/media"):
        self.docx_path = docx_path
        self.media_output_folder = media_output_folder
        self.doc = Document(docx_path)
        os.makedirs(self.media_output_folder, exist_ok=True)
        self.rel_images_map = self._mapear_relacionamentos_imagem()

    def _mapear_relacionamentos_imagem(self) -> Dict[str, str]:
        rel_map = {}
        try:
            with zipfile.ZipFile(self.docx_path, 'r') as z:
                for rel in self.doc.part.rels.values():
                    if "image" in rel.target_ref:
                        image_filename = os.path.basename(rel.target_ref)
                        zip_path = f"word/{rel.target_ref}".replace("//", "/")
                        if zip_path in z.namelist():
                            out_path = os.path.join(self.media_output_folder, image_filename)
                            with open(out_path, 'wb') as f:
                                f.write(z.read(zip_path))
                            rel_map[rel.rId] = out_path
        except Exception as e:
            print(f"[Aviso] Erro no mapeamento de imagens: {e}")
        return rel_map

    def extrair_imagens(self) -> List[str]:
        return list(self.rel_images_map.values())

    def _extrair_imagens_do_paragrafo(self, paragraph) -> List[str]:
        imgs = []
        element_xml = paragraph._element
        blips = element_xml.xpath('.//@r:embed')
        for rId in blips:
            if rId in self.rel_images_map:
                img_path = self.rel_images_map[rId]
                if img_path not in imgs:
                    imgs.append(img_path)
        return imgs

    def _extrair_formulas_omml(self, paragraph) -> str:
        formulas = []
        omath_elements = paragraph._element.xpath('.//m:oMath')
        for omath in omath_elements:
            texto_formula = "".join(omath.itertext()).strip()
            if omath.getparent().tag.endswith('oMathPara'):
                continue
            if texto_formula:
                formulas.append(texto_formula)
        return " ".join(formulas)

    def _normalizar_texto(self, texto: str) -> str:
        texto = texto.replace('\r', '\n').replace('\v', '\n').replace('\t', ' ')
        texto = re.sub(r' +', ' ', texto)
        return texto.strip()

    def _obter_linhas_do_documento(self) -> List[tuple]:
        linhas_processadas = []
        
        for p in self.doc.paragraphs:
            texto = self._normalizar_texto(p.text)
            imgs = self._extrair_imagens_do_paragrafo(p)
            formulas = self._extrair_formulas_omml(p)
            if formulas:
                texto = f"{texto} {formulas}".strip()
            if texto or imgs:
                linhas_processadas.append((p, texto, imgs))

        for table in self.doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        texto = self._normalizar_texto(p.text)
                        imgs = self._extrair_imagens_do_paragrafo(p)
                        formulas = self._extrair_formulas_omml(p)
                        if formulas:
                            texto = f"{texto} {formulas}".strip()
                        if texto or imgs:
                            linhas_processadas.append((p, texto, imgs))

        return linhas_processadas

    def processar(self) -> ProvaSchema:
        questoes: List[QuestaoSchema] = []
        questao_atual: Optional[QuestaoSchema] = None
        contador_questao = 0

        elementos = self._obter_linhas_do_documento()

        for paragraph, texto_bruto, imagens_paragrafo in elementos:
            if not texto_bruto and not imagens_paragrafo:
                continue

            sub_linhas = [l.strip() for l in texto_bruto.split('\n') if l.strip()] if texto_bruto else [""]

            for texto in sub_linhas:
                match_alt = self.REGEX_ALTERNATIVA.match(texto) if texto else None
                match_q = self.REGEX_NUMERO_QUESTAO.match(texto) if texto else None

                # 1. Processamento de Alternativas
                if match_alt:
                    letra = match_alt.group(1).lower()
                    texto_alt = match_alt.group(2).strip()

                    # Se encontrar 'a)' e a questão anterior já estava completa, abre uma nova questão sem lixo no enunciado
                    if not questao_atual or (letra == 'a' and len(questao_atual.alternativas) >= 4):
                        contador_questao += 1
                        questao_atual = QuestaoSchema(
                            numero=contador_questao,
                            enunciado="",
                            imagens=[],
                            alternativas=[]
                        )
                        questoes.append(questao_atual)

                    if questao_atual:
                        questao_atual.alternativas.append(
                            AlternativaSchema(letra=letra, texto=texto_alt)
                        )
                        if imagens_paragrafo:
                            questao_atual.imagens.extend(imagens_paragrafo)
                    continue

                # 2. Processamento de Título/Início de Questão
                eh_nova_questao = False
                if match_q:
                    eh_nova_questao = True
                elif not questao_atual:
                    eh_nova_questao = True

                if eh_nova_questao:
                    # Se a questão atual não tem alternativas e vem mais texto, concatena ao enunciado
                    if questao_atual and len(questao_atual.alternativas) == 0 and not match_q:
                        if texto:
                            if questao_atual.enunciado:
                                questao_atual.enunciado += f"\n{texto}"
                            else:
                                questao_atual.enunciado = texto
                        if imagens_paragrafo:
                            questao_atual.imagens.extend(imagens_paragrafo)
                        continue

                    contador_questao += 1
                    corpo = ""
                    if match_q:
                        corpo = match_q.group(2).strip() if match_q.group(2) else ""
                        if corpo.startswith("-") or corpo.startswith("–"):
                            corpo = corpo[1:].strip()
                    else:
                        corpo = texto

                    questao_atual = QuestaoSchema(
                        numero=contador_questao,
                        enunciado=corpo,
                        imagens=list(imagens_paragrafo),
                        alternativas=[]
                    )
                    questoes.append(questao_atual)
                    continue

                # 3. Anexo Limpo ao Enunciado
                if questao_atual:
                    if imagens_paragrafo:
                        for img in imagens_paragrafo:
                            if img not in questao_atual.imagens:
                                questao_atual.imagens.append(img)

                    if texto:
                        if len(questao_atual.alternativas) == 0:
                            if questao_atual.enunciado:
                                questao_atual.enunciado += f"\n{texto}"
                            else:
                                questao_atual.enunciado = texto
                        else:
                            # Texto secundário após alternativas é anexado limpo ao final do enunciado
                            questao_atual.enunciado += f"\n{texto}"

        return ProvaSchema(
            titulo="Avaliação do Colégio Santo Antônio",
            questoes=questoes
        )